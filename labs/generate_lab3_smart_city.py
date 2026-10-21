#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_centered(doc, text, size=14, bold=False, space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, bold=False, italic=False, first_line_indent=True, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)
    return p


def add_heading_like(doc, text, bold=True, size=14, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_figure_placeholder(doc, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("[Місце для рисунка]")
    set_run_font(run, size=12, italic=True)
    run.font.color.rgb = RGBColor(128, 128, 128)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    run = cap.add_run(caption)
    set_run_font(run, size=12, bold=True)
    return cap


def add_qa(doc, question, answer):
    add_body(doc, "Запитання:", bold=True, first_line_indent=False, space_after=2)
    add_body(doc, question, first_line_indent=False, space_after=4)
    add_body(doc, f"Відповідь: {answer}", first_line_indent=False, space_after=8)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_centered(doc, "Державний вищий навчальний заклад", size=14)
    add_centered(doc, "Ужгородський національний університет", size=14)
    add_centered(doc, "Факультет інформаційних технологій", size=14)
    add_centered(doc, "Кафедра програмного забезпечення систем", size=14, space_after=24)

    add_centered(doc, "ЛАБОРАТОРНА РОБОТА № 3", size=16, bold=True, space_before=24)
    add_centered(
        doc,
        "Тема: Packet Tracer – Дослідження розумного міста (Smart City)",
        size=14,
        bold=True,
        space_after=36,
    )

    add_centered(doc, "Виконала:", size=14, space_before=48)
    add_centered(doc, "студентка І курсу ОС Магістр", size=14)
    add_centered(doc, "спеціальності: 121 інженерія", size=14)
    add_centered(doc, "програмного забезпечення", size=14)
    add_centered(doc, "Чекета Юлія Михайлівна", size=14, bold=True, space_after=48)
    add_centered(doc, "Ужгород - 2026", size=14, space_before=72)

    doc.add_page_break()

    add_body(
        doc,
        "Мета роботи: Дослідити архітектуру розумного міста (Smart City) у Cisco Packet Tracer, "
        "проаналізувати роль хмарних і локальних IoT-серверів у моніторингу міських систем "
        "(паркування, трафік), а також спроєктувати та реалізувати IoT-пристрій «розумний стілець», "
        "що відстежує вагу та кількість посадок.",
    )

    add_heading_like(doc, "Цілі та задачі")
    add_body(doc, "Частина 1: Дослідження Smart City", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 2: Smart Parking", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 3: Smart Traffic", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 4: Рефлексія", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 5: Challenge — проєктування Smart Chair", first_line_indent=False, space_after=8)

    add_heading_like(doc, "Довідкова інформація / Сценарій")
    add_body(
        doc,
        "У цій лабораторній роботі досліджується приклад розумного міста. Ідея Smart City полягає "
        "в тому, що місто використовує програмне забезпечення для моніторингу та аналізу подій "
        "і прийняття рішень задля покращення життя мешканців. Місто розбито на кластери. "
        "Залежно від застосування частина даних обробляється та контролюється віддалено. "
        "Smart City Offices дають адміністратору можливість відстежувати події по всьому місту.",
    )

    # PART 1
    add_heading_like(doc, "Частина 1: Дослідження Smart City")
    add_heading_like(doc, "Крок 1: Пристрої, що входять до складу розумного міста", size=14)
    add_body(
        doc,
        "Оскільки датчики та IoT-пристрої розкидані містом, потрібна мережева інфраструктура. "
        "Smart City — це MAN (metropolitan area network), складений із менших мереж, часто "
        "з’єднаних WAN-каналами. ISP Cloud забезпечує загальноміський доступ.",
    )
    add_body(
        doc,
        "a. Відкрито ISP Cloud і переглянуто ресурси, які він надає місту.",
        first_line_indent=False,
    )
    add_figure_placeholder(doc, "Рис. 1 ISP Cloud та кластери Smart City")

    add_qa(
        doc,
        "Які міські мережі з’єднані червоними serial-кабелями?",
        "Smart Grid та City Offices.",
    )
    add_qa(
        doc,
        "Які міські мережі з’єднані синіми коаксіальними кабелями?",
        "Cell Tower0, City Offices та Smart Home.",
    )
    add_qa(
        doc,
        "Чому від ISP Cloud до City Offices ведуть два з’єднання?",
        "Одне з’єднання призначене для маршрутизації/загального доступу (router), "
        "друге — для підсистеми керування трафіком (traffic control).",
    )
    add_qa(
        doc,
        "Які міські мережі бездротово підключені до Cell-Tower?",
        "Smart Home, Smart Parking та City Offices.",
    )
    add_qa(
        doc,
        "Які пристрої в Smart Home підключені до Cell-Tower?",
        "Планшет (Tablet) та смартфон (Smartphone).",
    )
    add_qa(
        doc,
        "Які пристрої в кластері Smart Parking підключені до Cell-Tower?",
        "SmartPhone0 та S-Parking-RT.",
    )
    add_figure_placeholder(doc, "Рис. 2 Бездротові підключення до Cell-Tower")

    # PART 2
    add_heading_like(doc, "Частина 2: Smart Parking")
    add_body(
        doc,
        "Кластер Smart Parking дозволяє міській адміністрації та громадянам користуватися "
        "перевагами IoT-паркування.",
    )

    add_heading_like(doc, "Крок 1: Взаємодія з паркуванням (персонал City Offices)", size=14)
    add_body(
        doc,
        "Пристрої Smart Parking можна контролювати віддалено з будь-якого комп’ютера в City Offices, "
        "оскільки вони підключені до City IoT Server із вебінтерфейсом.",
        first_line_indent=False,
    )
    add_body(
        doc,
        "a–d. На City IT Laptop відкрито Desktop → Web Browser, введено адресу 195.0.0.2 "
        "(City IoT Server), виконано вхід Park / Park.",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Що відображається після входу на City IoT Server?",
        "Список зареєстрованих IoT-пристроїв паркування (лічильники/датчики паркувальних місць "
        "P-Space та пов’язані пристрої) з можливістю перегляду їхнього стану.",
    )
    add_figure_placeholder(doc, "Рис. 3 Вхід на City IoT Server (195.0.0.2), облікові дані Park/Park")

    add_qa(
        doc,
        "Яке значення відображається для P-Space-1 до постановки авто?",
        "Місце вільне (false / empty / 0 — місце не зайняте).",
    )
    add_body(
        doc,
        "g–h. Червоне авто перетягнуто на Parking Spot 1 (крайне ліве місце). "
        "Після оновлення на сервері стан P-Space-1 змінився.",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Яке значення відображається для P-Space-1 після постановки червоного авто?",
        "Місце зайняте (true / occupied / 1). Датчики місць — PT metal sensors, "
        "що реагують на металеві об’єкти (авто) поблизу.",
    )
    add_figure_placeholder(doc, "Рис. 4 P-Space-1: до і після постановки авто")

    add_heading_like(doc, "Крок 2: Взаємодія громадян (Smartphone)", size=14)
    add_body(
        doc,
        "Звичайні громадяни не повинні мати повний доступ до адмін-інтерфейсу сервера. "
        "Для них передбачено окрему вебсторінку локального parking-сервера на PT-MCU.",
    )
    add_body(
        doc,
        "b–c. На Smartphone відкрито Web Browser, введено адресу 10.10.10.10 "
        "(локальний parking server / PT-MCU).",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Що видно після завантаження сторінки на смартфоні?",
        "Вебсторінка доступності паркувальних місць: перелік місць і чи вони вільні/зайняті "
        "(громадський інтерфейс, без повного адмін-доступу).",
    )
    add_body(
        doc,
        "e–f. Зелене авто перетягнуто на Parking Spot 5 (крайнє праве). Сторінку на смартфоні оновлено.",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Що змінилось на сторінці після постановки зеленого авто на місце 5?",
        "Місце 5 відображається як зайняте / недоступне. Паркувальні місця повідомляють стан "
        "і на City IoT Server (City Offices), і на локальний parking server — громадяни можуть "
        "дізнатися про вільні місця ще до прибуття на вулицю.",
    )
    add_figure_placeholder(doc, "Рис. 5 Сторінка parking server 10.10.10.10 на смартфоні")
    add_body(
        doc,
        "h. Переглянуто код Parking Server MCU на вкладці Programming (за потреби — Advanced).",
        first_line_indent=False,
    )

    # PART 3
    add_heading_like(doc, "Частина 3: Smart Traffic")
    add_body(
        doc,
        "Smart Traffic дозволяє аварійним транспортним засобам (швидка, парамедики) "
        "спілкуватися зі світлофорною системою та запитувати вільний проїзд.",
    )
    add_body(
        doc,
        "a. Street Light 1 і Street Light 2 виконують роль світлофорів. "
        "Вантажівку парамедиків перетягнуто ближче до правого світлофора, який був червоним.",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Що відбувається зі світлофором справа?",
        "Світлофор стає зеленим, надаючи проїзд аварійному транспорту.",
    )
    add_figure_placeholder(doc, "Рис. 6 Швидка біля світлофора — зелений сигнал")

    add_body(
        doc,
        "c. Парамедиків відведено від правого світлофора та розміщено біля червоного авто.",
        first_line_indent=False,
    )
    add_qa(
        doc,
        "Що відбувається зі світлофором після від’їзду швидкої?",
        "Система повертається до звичайного режиму: після безпечного проїзду аварійного "
        "транспорту світлофори відновлюють нормальну роботу (правий знову може стати червоним).",
    )
    add_body(
        doc,
        "d. Транспорт парамедиків надсилає повідомлення MCU, що керує світлофорами, і запитує проїзд. "
        "MCU розпізнає його як легітимний аварійний транспорт, вмикає зелений і переводить "
        "пов’язані світлофори в червоний для безпеки. Після проїзду система повертається до норми.",
        first_line_indent=False,
    )
    add_body(
        doc,
        "e. Переглянуто код MCU в автомобілі швидкої на вкладці Programming (Advanced за потреби).",
        first_line_indent=False,
    )
    add_figure_placeholder(doc, "Рис. 7 Код MCU аварійного транспорту / світлофорів")

    # PART 4
    add_heading_like(doc, "Частина 4: Рефлексія")
    add_body(
        doc,
        "Коли датчики, програмне забезпечення та сервери взаємопов’язані, можна збирати дані "
        "й моніторити події. Рішення приймаються на основі подієвих даних. Ключовий висновок: "
        "«підключені речі» можуть спостерігати й діяти, якщо вони:",
    )
    add_body(doc, "• здатні виконувати програми;", first_line_indent=False, space_after=2)
    add_body(doc, "• підключені до датчиків або актуаторів;", first_line_indent=False, space_after=2)
    add_body(doc, "• з’єднані з іншими пристроями через мережу.", first_line_indent=False, space_after=6)
    add_body(
        doc,
        "Наприклад, звичайний стілець можна підготувати так, щоб він відстежував і записував "
        "вагу та кількість людей, які на нього сідають. Саме це реалізовано в Частині 5.",
    )

    # PART 5
    add_heading_like(doc, "Частина 5: Challenge — Smart Chair (опційно)")
    add_body(
        doc,
        "Завдання: спроєктувати розумний стілець, що відстежує та записує вагу людей, "
        "які на нього сідають.",
    )

    add_heading_like(doc, "Відповіді на керівні запитання", size=14)
    add_qa(
        doc,
        "Які деталі потрібні для проєкту (sensors, MCUs, SBCs, cables тощо)?",
        "MCU-PT (мозок), Potentiometer як модель датчика ваги (Force Sensor у цій версії PT "
        "відсутній), дві лампи LampLight і LampHeavy як індикатори, IoT Custom Cable для з’єднань.",
    )
    add_qa(
        doc,
        "Якою мовою писати програму?",
        "JavaScript у вкладці Programming мікроконтролера Packet Tracer.",
    )
    add_qa(
        doc,
        "Чи потрібні Internet-сервери або Cloud-підписки?",
        "Для мінімальної реалізації — ні. Вага й лічильник посадок виводяться в Serial Monitor MCU. "
        "За бажання можна додати Registration/IoT Server у PT без реальних хмарних підписок.",
    )
    add_qa(
        doc,
        "Чи можна побудувати блок-схеми виконання коду та роботи пристрою?",
        "Так. Нижче наведено логіку роботи та фрагмент реалізації.",
    )

    add_heading_like(doc, "Реалізація Smart Chair", size=14)
    add_body(doc, "Схема підключення:", bold=True, first_line_indent=False)
    add_body(doc, "• WeightSensor (Potentiometer) → MCU A0;", first_line_indent=False, space_after=2)
    add_body(doc, "• MCU D1 → LampLight;", first_line_indent=False, space_after=2)
    add_body(doc, "• MCU D2 → LampHeavy.", first_line_indent=False, space_after=6)
    add_body(
        doc,
        "Потенціометр імітує тензодатчик у сидінні: обертання моделює посадку/вставання людини. "
        "MCU зчитує analogRead(A0), переводить сигнал у діапазон 0–200 «кг», веде лічильник "
        "посадок (перехід empty → occupied) і керує лампами:",
    )
    add_body(doc, "• kg = 0 — empty, обидві лампи вимкнені;", first_line_indent=False, space_after=2)
    add_body(doc, "• 1 ≤ kg ≤ 50 — мигає LampLight;", first_line_indent=False, space_after=2)
    add_body(doc, "• kg ≥ 51 — мигає лише LampHeavy (LampLight вимкнено).", first_line_indent=False, space_after=6)
    add_figure_placeholder(doc, "Рис. 8 Топологія Smart Chair: MCU, WeightSensor, LampLight, LampHeavy")

    add_body(doc, "Фрагмент логіки програми (MCU, JavaScript):", bold=True, first_line_indent=False)
    add_body(
        doc,
        "var raw = analogRead(A0);\n"
        "var kg = raw * 200 / 1023;\n"
        "if (kg <= 0) { /* empty, лампи OFF */ }\n"
        "else if (kg <= 50) { /* мигає LampLight */ }\n"
        "else { /* мигає LampHeavy */ }",
        first_line_indent=False,
        space_after=8,
    )
    add_figure_placeholder(doc, "Рис. 9 Serial Monitor: kg, тип, NEW SIT, стан ламп")

    add_heading_like(doc, "Блок-схема роботи пристрою", size=14)
    add_body(
        doc,
        "Старт → setup (ініціалізація портів) → loop: читання A0 → обчислення kg → "
        "якщо нова посадка — totalSits++ → вибір режиму (empty / light / heavy) → "
        "керування лампами → delay → знову loop.",
        first_line_indent=False,
    )
    add_figure_placeholder(doc, "Рис. 10 Блок-схема алгоритму Smart Chair")

    add_heading_like(doc, "Висновок")
    add_body(
        doc,
        "Під час виконання лабораторної роботи досліджено топологію Smart City у Packet Tracer: "
        "роль ISP Cloud, serial- і coaxial-з’єднань, стільникового зв’язку через Cell-Tower. "
        "У кластері Smart Parking перевірено віддалений моніторинг через City IoT Server "
        "(195.0.0.2, Park/Park) та громадський інтерфейс локального parking server (10.10.10.10). "
        "У Smart Traffic продемонстровано пріоритет аварійного транспорту через взаємодію з MCU світлофорів. "
        "Це підтверджує використання хмарних/централізованих і локальних IoT-сервісів у місті. "
        "У challenge реалізовано Smart Chair на базі MCU, потенціометра та двох ламп: пристрій "
        "вимірює модельовану вагу, класифікує посадку, рахує кількість сідань і індикацією "
        "показує режим light/heavy — приклад перетворення звичайної речі на IoT-пристрій.",
    )

    out = "/workspace/labs/Лабораторна_робота_3_Smart_City.docx"
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
