#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_centered(doc, text, size=14, bold=False, space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, bold=False, italic=False, first_line_indent=True, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)
    return p


def add_heading_like(doc, text, bold=True, size=14, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_figure_placeholder(doc, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("[Місце для рисунка]")
    set_run_font(run, size=12, italic=True)
    run.font.color.rgb = RGBColor(128, 128, 128)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.space_before = Pt(0)
    run = cap.add_run(caption)
    set_run_font(run, size=12, bold=True)
    return cap


def add_qa(doc, question, answer):
    add_body(doc, "Запитання:", bold=True, first_line_indent=False, space_after=2)
    add_body(doc, question, first_line_indent=False, space_after=4)
    add_body(doc, f"Відповідь: {answer}", first_line_indent=False, space_after=8)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_centered(doc, "Державний вищий навчальний заклад", size=14)
    add_centered(doc, "Ужгородський національний університет", size=14)
    add_centered(doc, "Факультет інформаційних технологій", size=14)
    add_centered(doc, "Кафедра програмного забезпечення систем", size=14, space_after=24)

    add_centered(doc, "ЛАБОРАТОРНА РОБОТА № 2", size=16, bold=True, space_before=24)
    add_centered(
        doc,
        "Тема: Packet Tracer – Дослідження розумної енергомережі (Smart Grid)",
        size=14,
        bold=True,
        space_after=36,
    )

    add_centered(doc, "Виконала:", size=14, space_before=48)
    add_centered(doc, "студентка І курсу ОС Магістр", size=14)
    add_centered(doc, "спеціальності: 121 інженерія", size=14)
    add_centered(doc, "програмного забезпечення", size=14)
    add_centered(doc, "Чекета Юлія Михайлівна", size=14, bold=True, space_after=48)

    add_centered(doc, "Ужгород - 2026", size=14, space_before=72)

    doc.add_page_break()

    add_body(
        doc,
        "Мета роботи: Дослідити принципи побудови та функціонування розумної енергомережі (Smart Grid) "
        "у середовищі Cisco Packet Tracer, ознайомитися з пристроями генерації та моніторингу енергії, "
        "проаналізувати роботу програми Smart Power Grid Switch і роль мережевих пристроїв і хмарних/IoT-сервісів "
        "у виборі найефективнішого джерела живлення.",
    )

    add_heading_like(doc, "Цілі та задачі", bold=True)
    add_body(doc, "Частина 1: Дослідження Smart Grid", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 2: Рефлексія щодо використання обчислень і IoT у Smart Grid", first_line_indent=False, space_after=2)
    add_body(doc, "Частина 3: Додаткове завдання (Challenge Task)", first_line_indent=False, space_after=8)

    add_heading_like(doc, "Довідкова інформація / Сценарій", bold=True)
    add_body(
        doc,
        "У цій вправі досліджується розумне місто та розумна енергомережа (Smart Grid). "
        "Smart Grid відстежує потужність, що надходить з різних джерел, і перемикається на найефективніше "
        "доступне джерело живлення. Для цього використовується Smart Power Grid Switch.",
    )

    add_heading_like(doc, "Частина 1: Дослідження Smart Grid", bold=True)

    add_heading_like(doc, "Крок 1: Визначення пристроїв, що входять до складу Smart Grid", bold=True, size=14)
    add_body(
        doc,
        "Для роботи Smart Grid потрібні різні типи пристроїв. Мережеві пристрої забезпечують зв’язок усередині мережі. "
        "Сонячні панелі, вітрові турбіни та вугільна електростанція відповідають за генерацію енергії. "
        "Лічильники та датчики збирають дані й допомагають контролювати джерела живлення.",
    )

    add_qa(
        doc,
        "Перегляньте Smart Grid. Скільки маршрутизаторів ви бачите в Smart Grid? Як вони називаються?",
        "У мережі є 3 маршрутизатори: Wind_RT, Coal_RT та Main_RT (Power Main RT).",
    )
    add_figure_placeholder(doc, "Рис. 1 Топологія Smart Grid з маршрутизаторами Wind_RT, Coal_RT, Main_RT")

    add_body(
        doc,
        "Примітка: Packet Tracer підписує з’єднання між мережевими пристроями, але підписи можна вимкнути "
        "для зручності читання. Щоб увімкнути підписи портів: Options > Preferences > Interfaces tab > "
        "Always Show Port Labels in Logical Workspace.",
        italic=True,
    )

    add_qa(
        doc,
        "Яка функція маршрутизаторів?",
        "Маршрутизатори передають пакети даних між мережами (вітровою, вугільною та основною мережею Smart Grid), "
        "забезпечуючи зв’язок між сегментами енергомережі та доступ до Інтернету.",
    )

    add_qa(
        doc,
        "У Smart Grid використовується точка бездротового доступу. Як вона називається? Яка її функція?",
        "Назва: Power Access Point. Функція: забезпечує бездротове підключення та дозволяє "
        "віддалено керувати сонячними панелями (Solar Cells) за допомогою ноутбука.",
    )
    add_figure_placeholder(doc, "Рис. 2 Power Access Point у топології Smart Grid")

    add_qa(
        doc,
        "Чи можна визначити, яке джерело живлення активно виробляє енергію?",
        "Так. За допомогою лічильників потужності (Power Meter), установлених на кожному джерелі енергії "
        "(сонячні панелі, вітрова ферма, вугільна станція), можна побачити, яке джерело зараз активне.",
    )
    add_figure_placeholder(doc, "Рис. 3 Лічильники потужності на джерелах енергії")

    add_qa(
        doc,
        "Який пристрій відповідає за перемикання між різними джерелами живлення?",
        "Smart Power Grid Switch.",
    )

    add_qa(
        doc,
        "Як Smart Power Grid Switch вирішує, яке джерело живлення використовувати?",
        "Рішення приймає програма (JavaScript), запущена на перемикачі. Вона аналізує наявність енергії "
        "від кожного джерела. Пріоритет джерел: сонячна енергія > вітрова енергія > вугільна енергія. "
        "Якщо сонця немає, використовується вітер; якщо немає й вітру — вугільна станція.",
    )
    add_figure_placeholder(doc, "Рис. 4 Smart Power Grid Switch")

    add_qa(
        doc,
        "Яка IP-адреса Power IoT Server?",
        "100.2.0.2",
    )

    add_qa(
        doc,
        "Ноутбук використовується техніками для моніторингу кількості енергії. За допомогою веббраузера "
        "підключіться до Power IoT Server. Які ім’я користувача та пароль, і на якому пристрої їх виявлено?",
        "Ім’я користувача: Power; пароль: power. Облікові дані виявлено на пристрої Solar Cells "
        "(у налаштуваннях реєстрації на IoT-сервері).",
    )
    add_figure_placeholder(doc, "Рис. 5 Вхід на Power IoT Server (100.2.0.2) з ноутбука")

    add_heading_like(doc, "Крок 2: Дослідження програми Smart Power Grid Switch", bold=True, size=14)
    add_body(
        doc,
        "Smart Power Grid Switch може виконувати програми. У цьому прикладі програма написана мовою JavaScript. "
        "Вона відстежує рівні потужності від усіх джерел (вітрова ферма, сонячні панелі та вугільна станція) "
        "і визначає найкраще джерело.",
    )
    add_body(
        doc,
        "a. Клацнуто Smart Power Grid Switch.",
        first_line_indent=False,
    )
    add_body(
        doc,
        "b. Перейдено на вкладку Programming. Якщо вкладка не видима — натиснуто Advanced у нижній частині вікна.",
        first_line_indent=False,
    )
    add_body(
        doc,
        "c. Зліва двічі клацнуто power_switch(Javascript), потім main.js. "
        "Програма JavaScript відобразилася у правій частині вкладки Programming.",
        first_line_indent=False,
    )
    add_body(
        doc,
        "У місцевості, де розташована енергомережа, сонячне світло спостерігається з 6:00 до 18:00. "
        "Через близькість до океану вітер зазвичай дме з 3:00 до 22:00. Програма забезпечує пріоритет "
        "сонячної енергії. Якщо сонця немає — перемикається на вітер (якщо доступний). "
        "Якщо вітру немає — використовується вугільна енергія.",
    )
    add_figure_placeholder(doc, "Рис. 6 Програма main.js на вкладці Programming")

    add_qa(
        doc,
        "Який фрагмент програми робить цей пріоритет (solar > wind > coal) очевидним?",
        "Фрагмент з умовними операторами if / else if / else: спочатку перевіряється наявність сонячної енергії, "
        "потім вітрової, і лише якщо обидва джерела недоступні — обирається вугільна станція. "
        "Саме порядок перевірок у if-конструкції задає пріоритет solar > wind > coal.",
    )

    add_heading_like(doc, "Частина 2: Рефлексія", bold=True)
    add_body(
        doc,
        "Ключовий аспект IoT — інтелект, доданий до «речей». Програми дозволяють пристроям краще "
        "сприймати середовище й приймати рішення на основі подій і тригерів. Для цього можуть "
        "використовуватися різні мови програмування.",
    )
    add_body(
        doc,
        "Перш ніж програми зможуть виконуватися на «речі», потрібна відповідна апаратна платформа: "
        "по суті, перетворення пристрою на комп’ютер — шляхом додавання CPU, пам’яті та можливостей "
        "вводу/виводу (I/O).",
    )
    add_body(
        doc,
        "Нарешті, новостворена «розумна річ» повинна мати мережеві можливості зв’язку, "
        "щоб стати частиною IoT.",
    )
    add_body(
        doc,
        "У прикладі Smart Grid пристрою Smart Power Grid Switch надано обчислювальні можливості "
        "для запуску програм JavaScript. На основі датчиків і лічильників він усвідомлює стан середовища "
        "та діє згідно з наперед заданими умовами. Мережеві пристрої дозволяють перемикачу "
        "спілкуватися мережею й приєднуватися до IoT.",
    )

    add_heading_like(doc, "Частина 3: Додаткове завдання (Challenge Task)", bold=True)
    add_body(
        doc,
        "Завдання (опціонально): за допомогою Packet Tracer перетворити вікно на IoT-пристрій. "
        "Вікно має відстежувати рівень дощу в місцевості й при виявленні дощу зачинятися.",
    )
    add_body(doc, "Поради до виконання:", bold=True, first_line_indent=False)
    add_body(doc, "• Використати звичайне вікно PT (Window).", first_line_indent=False, space_after=2)
    add_body(
        doc,
        "• Призначити стани open/shut цифровому порту вікна: open = 1, shut = 0.",
        first_line_indent=False,
        space_after=2,
    )
    add_body(
        doc,
        "• На вкладці Programming відредагувати код для виявлення дощу та реакції на нього "
        "(за потреби натиснути Advanced).",
        first_line_indent=False,
        space_after=2,
    )
    add_body(
        doc,
        "• Вікно має запам’ятовувати останній стан (shut або open). Коли дощ закінчиться, "
        "вікно має повернутися до збереженого стану.",
        first_line_indent=False,
        space_after=2,
    )
    add_body(
        doc,
        "• Для зчитування значення дощу з симульованого середовища PT використати виклик "
        "Environment.get(\"Rain\"); (у деяких версіях інструкцій згадується Environment.get(\"Wind Speed\")).",
        first_line_indent=False,
        space_after=2,
    )
    add_figure_placeholder(doc, "Рис. 7 (опційно) IoT-вікно, що реагує на дощ")

    add_heading_like(doc, "Висновок", bold=True)
    add_body(
        doc,
        "Під час виконання даної лабораторної роботи я ознайомилася з топологією Smart Grid у Cisco Packet Tracer, "
        "визначила роль маршрутизаторів Wind_RT, Coal_RT і Main_RT, а також точки доступу Power Access Point. "
        "Було досліджено, як лічильники потужності показують активне джерело енергії, і як Smart Power Grid Switch "
        "за допомогою програми JavaScript обирає найкраще джерело за пріоритетом сонячна > вітрова > вугільна енергія. "
        "Через ноутбук перевірено віддалений доступ до Power IoT Server (IP 100.2.0.2) з обліковими даними Power/power, "
        "виявленими на пристрої Solar Cells. Це підтверджує, що «інтелект» IoT-пристроїв реалізується поєднанням "
        "апаратної платформи, програмної логіки та мережевого зв’язку.",
    )

    out = "/workspace/labs/Лабораторна_робота_2_Smart_Grid.docx"
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
